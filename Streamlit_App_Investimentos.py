# AVISO: Este aplicativo requer a instalação da biblioteca Streamlit.
# Para rodar localmente, certifique-se de que você executou:
# pip install streamlit

import streamlit as st
import pandas as pd
import sqlite3
import matplotlib.pyplot as plt
from datetime import datetime, timedelta
import numpy as np
import altair as alt
import openpyxl
import re
from pypfopt import EfficientFrontier, risk_models, expected_returns, plotting
from pypfopt.discrete_allocation import DiscreteAllocation, get_latest_prices
import plotly.graph_objects as go
from pypfopt.hierarchical_portfolio import HRPOpt
from st_aggrid import AgGrid, GridOptionsBuilder, JsCode

# Configuração da página
st.set_page_config(page_title="InvestApp", layout="wide")

st.write("🔄 Aplicativo carregado. Tentando carregar os dados...")

# --- Conectar ao banco de dados SQLite ---
DB_PATH = "dados_investimentos.db"

# --- Funções auxiliares ---
def carregar_dados():
    conn = sqlite3.connect(DB_PATH)
    df = pd.read_sql("SELECT date, ticker, close FROM dados_historicos", conn)
    conn.close()
    df['date'] = pd.to_datetime(df['date'])
    return df

def salvar_na_carteira(nome, dados):
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    cursor.execute('''CREATE TABLE IF NOT EXISTS carteiras (
                        id INTEGER PRIMARY KEY AUTOINCREMENT,
                        nome TEXT,
                        ticker TEXT,
                        quantidade INTEGER,
                        valor_compra REAL,
                        data_compra TEXT
                    );''')
    for linha in dados:
        cursor.execute("INSERT INTO carteiras (nome, ticker, quantidade, valor_compra, data_compra) VALUES (?, ?, ?, ?, ?);",
                       (nome, linha['ticker'], linha['quantidade'], linha['valor_compra'], linha['data_compra']))
    conn.commit()
    conn.close()

def carregar_carteiras_disponiveis():
    conn = sqlite3.connect(DB_PATH)
    df = pd.read_sql("SELECT DISTINCT nome FROM carteiras", conn)
    conn.close()
    return df['nome'].tolist()

def carregar_carteira(nome):
    conn = sqlite3.connect(DB_PATH)
    df = pd.read_sql("SELECT id, ticker, quantidade, valor_compra, data_compra FROM carteiras WHERE nome = ?", conn, params=(nome,))
    conn.close()
    return df

def atualizar_entrada_carteira(id, quantidade, valor, data):
    conn = sqlite3.connect(DB_PATH)
    conn.execute("UPDATE carteiras SET quantidade = ?, valor_compra = ?, data_compra = ? WHERE id = ?",
                 (quantidade, valor, data, id))
    conn.commit()
    conn.close()


def importar_tabela_excel(uploaded_file):
    wb = openpyxl.load_workbook(uploaded_file, data_only=True)
    
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        if sheet.tables:
            tabela = list(sheet.tables.values())[0]
            ref = tabela.ref

            col_range = ":".join(re.findall(r"[A-Z]+", ref))
            
            df = pd.read_excel(uploaded_file, sheet_name=sheet_name, usecols=col_range)
            df.columns = [col.lower().strip() for col in df.columns]
            return df

    raise ValueError("Nenhuma tabela nomeada foi encontrada no Excel.")


def simular_carteira(df_historico, ativos):
    registros_diarios = []

    for aporte in ativos:
        ticker = aporte['ticker']
        qtd = aporte['quantidade']
        preco_compra = aporte['valor_compra']
        data_ini = pd.to_datetime(aporte['data_compra'])

        df_ativo = df_historico[(df_historico['ticker'] == ticker) & (df_historico['date'] >= data_ini)].copy()
        df_ativo = df_ativo.sort_values('date')
        if df_ativo.empty:
            continue

        df_ativo['valor_mercado'] = df_ativo['close'] * qtd
        df_ativo['aporte'] = preco_compra * qtd
        df_ativo['ticker'] = ticker
        df_ativo['id_aporte'] = f"{ticker}_{data_ini.strftime('%Y%m%d')}"
        df_ativo['data_aporte'] = data_ini
        registros_diarios.append(df_ativo[['date', 'valor_mercado', 'aporte', 'id_aporte', 'data_aporte']])

    if not registros_diarios:
        return pd.DataFrame()

    df_todos = pd.concat(registros_diarios)

    # Soma dos valores de mercado por dia
    df_valor_diario = df_todos.groupby('date')['valor_mercado'].sum().reset_index(name='valor_total')

    # Soma acumulada de aportes até cada dia
    df_aportes = df_todos[['id_aporte', 'data_aporte', 'aporte']].drop_duplicates()
    df_aportes = df_aportes.sort_values('data_aporte')

    df_valor_diario['aporte_acumulado'] = df_valor_diario['date'].apply(
        lambda d: df_aportes[df_aportes['data_aporte'] <= d]['aporte'].sum()
    )

    # Proteção contra divisão por zero
    df_valor_diario['rendimento_acumulado'] = df_valor_diario.apply(
        lambda row: row['valor_total'] / row['aporte_acumulado'] - 1 if row['aporte_acumulado'] > 0 else 0,
        axis=1
    )

    return df_valor_diario


def carregar_detalhes(ticker_input):
    conn = sqlite3.connect(DB_PATH)

    # Histórico mais recente do ativo
    query_hist = f"""
        SELECT * FROM dados_historicos
        WHERE ticker = '{ticker_input}'
        ORDER BY date DESC
        LIMIT 1
    """
    df_hist = pd.read_sql(query_hist, conn)
    hist = df_hist.iloc[0] if not df_hist.empty else {}

    # Indicadores fundamentais
    query_fund = f"SELECT * FROM dados_fundamentais WHERE ticker = '{ticker_input}'"
    df_fund = pd.read_sql(query_fund, conn)
    fund = df_fund.iloc[0] if not df_fund.empty else {}

    # Indicadores macro (CDI, IPCA)
    df_ind = pd.read_sql("SELECT * FROM dados_indicadores", conn)
    df_ind['data'] = pd.to_datetime(df_ind['data'])

    conn.close()
    return hist, fund, df_ind

@st.cache_data
def calcular_tabela_consolidada():
    conn = sqlite3.connect(DB_PATH)
    df_hist = pd.read_sql("SELECT date, ticker, close FROM dados_historicos", conn)
    df_fund = pd.read_sql("SELECT ticker, nome, setor FROM dados_fundamentais", conn)
    conn.close()

    df_hist['date'] = pd.to_datetime(df_hist['date'])
    ultimo_dia = df_hist['date'].max()
    data_1ano = ultimo_dia - timedelta(days=365)
    data_30d = ultimo_dia - timedelta(days=30)

    resultados = []
    for ticker in df_hist['ticker'].unique():
        df_ticker = df_hist[df_hist['ticker'] == ticker].sort_values('date')
        df_1ano = df_ticker[df_ticker['date'] >= data_1ano]
        df_30d = df_ticker[df_ticker['date'] >= data_30d]
        if df_1ano.empty or df_30d.empty:
            continue
        preco_atual = df_ticker['close'].iloc[-1]
        variacao_ano = preco_atual / df_1ano['close'].iloc[0] - 1
        variacao_30d = preco_atual / df_30d['close'].iloc[0] - 1
        media_movel_ano = df_1ano['close'].mean()
        indice_mayer = preco_atual / media_movel_ano
        resultados.append({
            'Ticker': ticker,
            'Variação 1 ano (%)': variacao_ano * 100,
            'Variação 30 dias (%)': variacao_30d * 100,
            'Índice Mayer': indice_mayer
        })

    df_resultados = pd.DataFrame(resultados)
    df_resultados = df_resultados.merge(df_fund[['ticker', 'nome', 'setor']], left_on='Ticker', right_on='ticker', how='left')
    df_resultados.drop(columns='ticker', inplace=True)
    df_resultados.rename(columns={'nome': 'Nome da Empresa', 'setor': 'Setor'}, inplace=True)
    return df_resultados

@st.cache_data
def calcular_variacao_por_setor(df_filtrado):
    return (
        df_filtrado
        .groupby('Setor')[['Variação 1 ano (%)', 'Variação 30 dias (%)']]
        .mean()
        .reset_index()
    )

def calcular_performances(df_historico, tickers, data_fim):
    """Calcula a performance de uma lista de tickers para vários períodos."""
    performances = {}
    df_historico = df_historico.sort_values('date')

    data_1ano = data_fim - timedelta(days=365)
    data_30d = data_fim - timedelta(days=30)
    data_ytd = pd.to_datetime(f"{data_fim.year}-01-01")

    for ticker in tickers:
        df_ticker = df_historico[df_historico['ticker'] == ticker]
        if df_ticker.empty:
            performances[ticker] = {'30d': 0, '1y': 0, 'ytd': 0}
            continue

        preco_atual_row = df_ticker[df_ticker['date'] <= data_fim]
        if preco_atual_row.empty:
            performances[ticker] = {'30d': 0, '1y': 0, 'ytd': 0}
            continue
        preco_atual = preco_atual_row['close'].iloc[-1]

        def get_start_price(df, start_date):
            rows = df[df['date'] >= start_date]
            return rows['close'].iloc[0] if not rows.empty else None

        preco_30d_atras = get_start_price(df_ticker, data_30d)
        perf_30d = (preco_atual / preco_30d_atras - 1) if preco_30d_atras else 0

        preco_1ano_atras = get_start_price(df_ticker, data_1ano)
        perf_1y = (preco_atual / preco_1ano_atras - 1) if preco_1ano_atras else 0

        preco_ytd_inicio = get_start_price(df_ticker, data_ytd)
        perf_ytd = (preco_atual / preco_ytd_inicio - 1) if preco_ytd_inicio else 0

        performances[ticker] = {'30d': perf_30d, '1y': perf_1y, 'ytd': perf_ytd}
    return performances


def rebalance_weights_with_cap(weights, max_cap):

    rebalanced_weights = weights.copy()
    
    while True:
        overweight_tickers = {t for t, w in rebalanced_weights.items() if w > max_cap}
        
        if not overweight_tickers:
            break 
            
        total_excess = sum(rebalanced_weights[t] - max_cap for t in overweight_tickers)
        
        for t in overweight_tickers:
            rebalanced_weights[t] = max_cap
        
        underweight_tickers = {t for t, w in rebalanced_weights.items() if w < max_cap}
        
        if not underweight_tickers:
            break
            
        sum_of_underweight = sum(rebalanced_weights[t] for t in underweight_tickers)
        
        if sum_of_underweight > 0:
            for t in underweight_tickers:
                rebalanced_weights[t] += total_excess * (rebalanced_weights[t] / sum_of_underweight)
        
    total_weight = sum(rebalanced_weights.values())
    for t in rebalanced_weights:
        rebalanced_weights[t] /= total_weight
        
    return rebalanced_weights

# --- FUNÇÕES DE BACKTEST ---
def get_weights(method, returns_df, max_alloc):
    weights = {}
    try:
        if method == "HRP":
            hrp = HRPOpt(returns_df)
            hrp.optimize()
            weights = hrp.clean_weights()
        elif method == "Min Volatility":
            S = risk_models.sample_cov(returns_df, returns_data=True)
            if not S.is_positive_semidefinite():
                S = risk_models.fix_non_positive_semidefinite(S, fix_method='spectral')
            ef = EfficientFrontier(None, S)
            ef.min_volatility()
            weights = ef.clean_weights()
        elif method in ["Equal Weight", "Rebalancear para Pesos Iguais"]:
            num_assets = len(returns_df.columns)
            weights = {ticker: 1/num_assets for ticker in returns_df.columns}
        else:
            raise ValueError("Método de alocação desconhecido.")
        
        if max_alloc < 1.0 and method != "Rebalancear para Pesos Iguais":
            weights = rebalance_weights_with_cap(weights, max_alloc)
    except Exception:
        num_assets = len(returns_df.columns)
        weights = {ticker: 1/num_assets for ticker in returns_df.columns}
    return weights

@st.cache_data(show_spinner=False)
def run_backtest(_df_historico, tickers, start_date, end_date, initial_investment, monthly_contribution, max_alloc, method):
    df_prices = _df_historico[(_df_historico['ticker'].isin(tickers)) & (_df_historico['date'] >= start_date - timedelta(days=365*2))].copy()
    df_prices_pivot = df_prices.pivot(index='date', columns='ticker', values='close')

    rebalance_dates = pd.date_range(start_date, end_date, freq='BMS')
    holdings = {ticker: 0.0 for ticker in tickers}
    total_invested_per_asset = {ticker: 0.0 for ticker in tickers}
    daily_history_list = []
    capital_invested = 0.0
    weights = {}
    asset_start_dates = {}
    actual_start_date = None
    first_investment_made = False

    for i, date in enumerate(rebalance_dates):
        try:
            current_rebalance_date = df_prices_pivot.index[df_prices_pivot.index.searchsorted(date)]
        except IndexError:
            continue
        
        end_period = rebalance_dates[i+1] if i + 1 < len(rebalance_dates) else end_date
        
        if first_investment_made:
            last_prices = df_prices_pivot.loc[:current_rebalance_date].iloc[-1]
            portfolio_value = sum(shares * last_prices.get(ticker, 0) for ticker, shares in holdings.items() if shares > 0)
        else:
            portfolio_value = 0
            
        data_slice = df_prices_pivot.loc[:current_rebalance_date]
        hist_data_for_opt = data_slice.dropna(axis=1, how='all')
        
        if hist_data_for_opt.shape[1] < 2 or len(hist_data_for_opt.dropna()) < 60:
            if first_investment_made:
                daily_prices_slice = df_prices_pivot.loc[current_rebalance_date:end_period]
                if not daily_prices_slice.empty:
                    daily_values = daily_prices_slice.multiply(pd.Series(holdings), axis='columns').sum(axis=1)
                    df_period_history = pd.DataFrame({'value': daily_values, 'invested': capital_invested})
                    daily_history_list.append(df_period_history)
            continue

        is_first_rebalance = not first_investment_made
        if is_first_rebalance:
            actual_start_date = current_rebalance_date
            first_investment_made = True
            capital_invested = initial_investment
            total_capital = initial_investment
        else:
            capital_invested += monthly_contribution
            total_capital = portfolio_value + monthly_contribution
        
        returns = hist_data_for_opt.pct_change().dropna()
        returns = returns.loc[:, returns.std() > 1e-8]
        if returns.shape[1] < 2:
            if first_investment_made:
                daily_prices_slice = df_prices_pivot.loc[current_rebalance_date:end_period]
                if not daily_prices_slice.empty:
                    daily_values = daily_prices_slice.multiply(pd.Series(holdings), axis='columns').sum(axis=1)
                    df_period_history = pd.DataFrame({'value': daily_values, 'invested': capital_invested})
                    daily_history_list.append(df_period_history)
            continue
        
        weights = get_weights(method, returns, max_alloc)
        for ticker in weights:
            if ticker not in asset_start_dates:
                asset_start_dates[ticker] = current_rebalance_date

        latest_prices = hist_data_for_opt.iloc[-1]
        
        if method == "Rebalancear para Pesos Iguais" and not is_first_rebalance:
            aporte_externo_neste_mes = monthly_contribution
            current_weights = {t: (holdings[t] * latest_prices.get(t,0)) / portfolio_value for t in holdings} if portfolio_value > 0 else {}
            underweight_assets = {t: weights[t] - current_weights.get(t, 0) for t in weights if current_weights.get(t, 0) < weights[t]}
            
            if underweight_assets:
                total_shortfall = sum(underweight_assets.values())
                for ticker, shortfall in underweight_assets.items():
                    aporte_para_ativo = aporte_externo_neste_mes * (shortfall / total_shortfall)
                    total_invested_per_asset[ticker] += aporte_para_ativo
                    holdings[ticker] += aporte_para_ativo / latest_prices.get(ticker, 1)
        else:
            aporte_externo_neste_mes = initial_investment if is_first_rebalance else monthly_contribution
            for ticker_alloc, weight in weights.items():
                total_invested_per_asset[ticker_alloc] += aporte_externo_neste_mes * weight
            for ticker_alloc in tickers:
                if ticker_alloc in weights and latest_prices.get(ticker_alloc, 0) > 0:
                    valor_alocado = total_capital * weights.get(ticker_alloc, 0)
                    holdings[ticker_alloc] = valor_alocado / latest_prices.get(ticker_alloc, 0)
                else:
                    holdings[ticker_alloc] = 0.0
        
        daily_prices_slice = df_prices_pivot.loc[current_rebalance_date:end_period]
        if not daily_prices_slice.empty:
            daily_values = daily_prices_slice.multiply(pd.Series(holdings), axis='columns').sum(axis=1)
            df_period_history = pd.DataFrame({'value': daily_values, 'invested': capital_invested})
            daily_history_list.append(df_period_history)

    if not daily_history_list:
        return pd.DataFrame(), {}

    df_portfolio = pd.concat(daily_history_list)
    df_portfolio = df_portfolio[~df_portfolio.index.duplicated(keep='last')]
            
    final_state = { "holdings": holdings, "prices": df_prices_pivot.iloc[-1] if not df_prices_pivot.empty else pd.Series(), "weights": weights, "start_dates": asset_start_dates, "actual_start_date": actual_start_date, "total_invested_per_asset": total_invested_per_asset }
    return df_portfolio, final_state

def get_benchmark_data(start_date, end_date):
    if start_date is None:
        return pd.DataFrame(columns=['data', 'IPCA', 'CDI', 'IBOV'])

    conn = sqlite3.connect(DB_PATH)
    df_indicadores = pd.read_sql("SELECT * FROM dados_indicadores", conn)
    conn.close()
    df_indicadores['data'] = pd.to_datetime(df_indicadores['data'])

    df_ind_periodo = df_indicadores[(df_indicadores['data'] >= start_date) & (df_indicadores['data'] <= end_date)].copy().sort_values('data')
    if df_ind_periodo.empty:
        return pd.DataFrame(columns=['data', 'IPCA', 'CDI', 'IBOV'])

    df_ind_periodo['ipca_norm'] = (1 + df_ind_periodo['ipca'].fillna(0))
    df_ind_periodo['IPCA'] = df_ind_periodo['ipca_norm'].cumprod() / df_ind_periodo['ipca_norm'].iloc[0] - 1
    
    df_ind_periodo['cdi_norm'] = (1 + df_ind_periodo['cdi'].fillna(0))
    df_ind_periodo['CDI'] = df_ind_periodo['cdi_norm'].cumprod() / df_ind_periodo['cdi_norm'].iloc[0] - 1
    
    if 'ibov' in df_ind_periodo.columns and df_ind_periodo['ibov'].notna().sum() >= 2:
        ibov_validos = df_ind_periodo['ibov'].dropna()
        primeiro_ibov = ibov_validos.iloc[0]
        df_ind_periodo['IBOV'] = df_ind_periodo['ibov'].apply(lambda x: x / primeiro_ibov - 1 if pd.notna(x) else np.nan)
        df_ind_periodo['IBOV'] = df_ind_periodo['IBOV'].ffill()
    else:
        df_ind_periodo['IBOV'] = np.nan
        
    return df_ind_periodo[['data', 'IPCA', 'CDI', 'IBOV']]


# --- Tentativa segura de carregar os dados ---
try:
    df = carregar_dados()
    tickers_unicos = sorted(df['ticker'].unique())
    st.success("✅ Dados carregados com sucesso.")
except Exception as e:
    st.error(f"❌ Erro ao carregar os dados: {e}")
    st.stop()

# --- Função para calcular rendimento acumulado ---
def calcular_rendimento(df, ticker, data_inicio, data_fim):
    df_filtrado = df[(df['ticker'] == ticker) & 
                     (df['date'] >= data_inicio) & 
                     (df['date'] <= data_fim)].copy()
    df_filtrado.sort_values('date', inplace=True)
    df_filtrado['rendimento_acumulado'] = df_filtrado['close'] / df_filtrado['close'].iloc[0] - 1
    return df_filtrado

# --- Menu lateral ---
st.sidebar.title("📊 Menu de Navegação")
opcao = st.sidebar.selectbox("Escolha uma opção:", [
    "📈 Visão por Ativo",
    "📊 Visão Geral",
    "📊 Análise Carteira",
    "🤖 Otimizador Avançado (HRP)",
    "🧪 Backtesting de Estratégias",
    "📉 Detecção de Rompimentos Técnicos",
    "📂 Carteira de Investimento"
], key="opcao")

# --- Visão por Ativo ---
if opcao == "📈 Visão por Ativo":
    st.title("📈 Visão por Ativo")

    params = st.query_params
    if "ticker" in params:
        st.session_state["ticker_input"] = params["ticker"]
        st.query_params.clear()

    ticker_default = (
        st.session_state["ticker_input"]
        if "ticker_input" in st.session_state and st.session_state["ticker_input"] in tickers_unicos
        else tickers_unicos[0]
    )

    ticker_input = st.selectbox("Selecione o ativo:", tickers_unicos, index=tickers_unicos.index(ticker_default))

    df_ativo = df[df['ticker'] == ticker_input].copy()
    df_ativo.sort_values('date', inplace=True)

    if df_ativo.empty:
        st.warning("Não há dados disponíveis para este ativo.")
        st.stop()

    hoje = df_ativo['date'].max()
    datas_opcoes = {
        "30 dias": hoje - timedelta(days=30),
        "6 meses": hoje - timedelta(days=182),
        "1 ano": hoje - timedelta(days=365),
        "5 anos": hoje - timedelta(days=1825),
        "10 anos": hoje - timedelta(days=3650),
        "YTD": pd.to_datetime(f"{hoje.year}-01-01"),
        "Total": df_ativo['date'].min()
    }

    selecao_periodo = st.radio("Período:", list(datas_opcoes.keys()), horizontal=True)
    data_inicio = datas_opcoes[selecao_periodo]

    df_periodo = df_ativo[df_ativo['date'] >= data_inicio].copy()
    df_periodo['rendimento'] = df_periodo['close'] / df_periodo['close'].iloc[0] - 1
    variacao_percentual = df_periodo['rendimento'].iloc[-1] * 100

    hist, fund, df_ind = carregar_detalhes(ticker_input)

    df_ind_periodo = df_ind[(df_ind['data'] >= data_inicio) & (df_ind['data'] <= hoje)].copy()
    variacao_cdi = (1 + df_ind_periodo['cdi'].fillna(0)).prod() - 1 if not df_ind_periodo.empty else 0
    variacao_ipca = (1 + df_ind_periodo['ipca'].fillna(0)).prod() - 1 if not df_ind_periodo.empty else 0
    if 'ibov' in df_ind_periodo.columns:
        ibov_validos = df_ind_periodo['ibov'].dropna()
        if not ibov_validos.empty and ibov_validos.index[0] != ibov_validos.index[-1]:
            primeiro_valido = ibov_validos.iloc[0]
            ultimo_valido = ibov_validos.iloc[-1]
            variacao_ibov = ultimo_valido / primeiro_valido - 1
        else:
            variacao_ibov = 0
    else:
        variacao_ibov = 0


    cotacao = f"R$ {df_ativo['close'].iloc[-1]:.2f}"
    variacao_percentual_fmt = f"{variacao_percentual:.1f}%"
    variacao_cdi_fmt = f"{variacao_cdi * 100:.1f}%"
    variacao_ipca_fmt = f"{variacao_ipca * 100:.1f}%"
    variacao_ibov_fmt = f"{variacao_ibov * 100:.1f}%"
    cor_variacao = 'green' if variacao_percentual >= 0 else 'red'

    valores = {
        "variacao": variacao_percentual,
        "cdi": variacao_cdi * 100,
        "ipca": variacao_ipca * 100,
        "ibov": variacao_ibov * 100
    }
    maior_chave = max(valores, key=valores.get)

    col0, col1, col2, col3, col4 = st.columns(5)
    col0.markdown(f"<div style='text-align:center'><div style='font-size:18px; font-weight:bold;'>Cotação</div><p style='font-size: 24px;'>{cotacao}</p></div>", unsafe_allow_html=True)
    col1.markdown(f"<div style='text-align:center'><div style='font-size:18px; font-weight:bold;'>Variação ({selecao_periodo})</div><p style='color:{cor_variacao}; font-size: 24px;'>{variacao_percentual_fmt}{' ⭐' if maior_chave == 'variacao' else ''}</p></div>", unsafe_allow_html=True)
    col2.markdown(f"<div style='text-align:center'><div style='font-size:18px; font-weight:bold;'>CDI</div><p style='font-size: 24px;'>{variacao_cdi_fmt}{' ⭐' if maior_chave == 'cdi' else ''}</p></div>", unsafe_allow_html=True)
    col3.markdown(f"<div style='text-align:center'><div style='font-size:18px; font-weight:bold;'>IPCA</div><p style='font-size: 24px;'>{variacao_ipca_fmt}{' ⭐' if maior_chave == 'ipca' else ''}</p></div>", unsafe_allow_html=True)
    col4.markdown(f"<div style='text-align:center'><div style='font-size:18px; font-weight:bold;'>IBOV</div><p style='font-size: 24px;'>{variacao_ibov_fmt}{' ⭐' if maior_chave == 'ibov' else ''}</p></div>", unsafe_allow_html=True)

    df_periodo['timestamp'] = df_periodo['date'].map(pd.Timestamp.timestamp)
    x = df_periodo['timestamp']
    y = df_periodo['close']
    coef = np.polyfit(x, y, 1)
    trend = np.poly1d(coef)
    df_periodo['trend'] = trend(x)
    residuo = df_periodo['close'] - df_periodo['trend']
    std = residuo.std()
    df_periodo['upper_outlier'] = df_periodo['trend'] + std
    df_periodo['lower_outlier'] = df_periodo['trend'] - std

    mostrar_linhas = st.checkbox("📉 Exibir linha de tendência e outliers", value=True)

    if not mostrar_linhas:
        chart = alt.Chart(df_periodo).mark_line(color='deepskyblue').encode(
            x=alt.X("date:T", title="Data", axis=alt.Axis(format="%d/%m/%Y", labelAngle=-45)),
            y=alt.Y("close:Q", title="Preço (R$)", scale=alt.Scale(zero=False)),
            tooltip=["date:T", alt.Tooltip("close:Q", title="Preço")]
        ).properties(width=800, height=400)
    else:
        df_long = df_periodo[['date', 'close', 'trend', 'upper_outlier', 'lower_outlier']].copy()
        df_long = df_long.rename(columns={
            'close': 'Preço',
            'trend': 'Tendência',
            'upper_outlier': 'Outlier Superior',
            'lower_outlier': 'Outlier Inferior'
        }).melt(id_vars='date', var_name='Indicador', value_name='Valor')

        base = alt.Chart(df_long).encode(
            x=alt.X("date:T", title="Data"),
            y=alt.Y("Valor:Q", title="Preço (R$)", scale=alt.Scale(zero=False))
        )


        highlight = alt.selection_point(
            nearest=True, on="mouseover", fields=["date"], empty="none", clear="mouseout"
        )

        linha_preco = base.transform_filter(alt.datum.Indicador == "Preço").mark_line(color="deepskyblue")
        linha_tendencia = base.transform_filter(alt.datum.Indicador == "Tendência").mark_line(color="red")
        linha_outlier_sup = base.transform_filter(alt.datum.Indicador == "Outlier Superior").mark_line(
            color="gray", strokeDash=[4, 4]
        )
        linha_outlier_inf = base.transform_filter(alt.datum.Indicador == "Outlier Inferior").mark_line(
            color="gray", strokeDash=[4, 4]
        )

        pontos = base.mark_circle(size=60).encode(opacity=alt.value(0)).add_params(highlight)
        texto = base.mark_text(align="left", dx=5, dy=-5).encode(
            text=alt.condition(highlight, alt.Text("Valor:Q", format=".2f"), alt.value(""))
        )
        regua = alt.Chart(df_long).mark_rule(color="gray").encode(x="date:T").transform_filter(highlight)

        chart = alt.layer(
            linha_preco, linha_tendencia, linha_outlier_sup, linha_outlier_inf, pontos, texto, regua
        ).properties(width=800, height=400)


    st.altair_chart(chart, use_container_width=True)

    st.markdown("### 📊 Indicadores Fundamentais")

    pl = f"{hist['trailing_pe']:.2f}" if 'trailing_pe' in hist and pd.notna(hist['trailing_pe']) else "-"
    pvp = f"{hist['price_to_book']:.2f}" if 'price_to_book' in hist and pd.notna(hist['price_to_book']) else "-"
    roe = f"{fund['roe'] * 100:.2f}%" if 'roe' in fund and pd.notna(fund['roe']) else "-"

    col_a, col_b, col_c = st.columns(3)
    col_a.markdown(f"<div style='text-align:center'><h4>P/L</h4><p style='font-size: 24px;'>{pl}</p></div>", unsafe_allow_html=True)
    col_b.markdown(f"<div style='text-align:center'><h4>P/VP</h4><p style='font-size: 24px;'>{pvp}</p></div>", unsafe_allow_html=True)
    col_c.markdown(f"<div style='text-align:center'><h4>ROE</h4><p style='font-size: 24px;'>{roe}</p></div>", unsafe_allow_html=True)

# --- Visão Geral ---
elif opcao == "📊 Visão Geral":
    st.title("📊 Visão Geral Consolidada")

    df_resultados = calcular_tabela_consolidada()

    col1, col2 = st.columns([1, 1])
    setores = sorted(df_resultados['Setor'].dropna().unique())
    with col1:
        setor_escolhido = st.selectbox("Setor", options=[""] + setores, index=0)

    mayer_min = df_resultados['Índice Mayer'].min()
    mayer_max = df_resultados['Índice Mayer'].max()
    with col2:
        filtro_mayer = st.slider("Filtrar Índice Mayer", min_value=float(mayer_min), max_value=float(mayer_max),
                                 value=(float(mayer_min), float(mayer_max)))

    df_filtrado = df_resultados[
        (df_resultados['Índice Mayer'] >= filtro_mayer[0]) & 
        (df_resultados['Índice Mayer'] <= filtro_mayer[1])
    ]
    if setor_escolhido:
        df_filtrado = df_filtrado[df_filtrado['Setor'] == setor_escolhido]

    df_setor_var = calcular_variacao_por_setor(df_filtrado)

    col_a, col_b, col_c = st.columns(3)
    if col_a.button("Classificar por Variação 1 ano"):
        st.session_state["ordenar_por"] = "Variação 1 ano (%)"
    if col_b.button("Classificar por Variação 30 dias"):
        st.session_state["ordenar_por"] = "Variação 30 dias (%)"
    if col_c.button("Classificar por Índice Mayer"):
        st.session_state["ordenar_por"] = "Índice Mayer"

    ordenar_por = st.session_state.get("ordenar_por", None)
    if ordenar_por:
        df_filtrado = df_filtrado.sort_values(by=ordenar_por, ascending=False)
        if ordenar_por in df_setor_var.columns:
            df_setor_var = df_setor_var.sort_values(by=ordenar_por, ascending=False)

    if not df_setor_var.empty:
        html_setores = "<table style='width:100%; border-collapse: collapse;'>"
        html_setores += "<thead><tr>"
        html_setores += "<th style='text-align:center; border-bottom: 1px solid #ddd'>Setor</th>"
        html_setores += "<th style='text-align:center; border-bottom: 1px solid #ddd'>Variação 1 ano (%)</th>"
        html_setores += "<th style='text-align:center; border-bottom: 1px solid #ddd'>Variação 30 dias (%)</th>"
        html_setores += "</tr></thead><tbody>"

        for _, row in df_setor_var.iterrows():
            cor_ano = "green" if row["Variação 1 ano (%)"] >= 1 else "red"
            cor_30d = "green" if row["Variação 30 dias (%)"] >= 1 else "red"

            html_setores += "<tr>"
            html_setores += f"<td style='text-align:center'>{row['Setor']}</td>"
            html_setores += f"<td style='color:{cor_ano}; text-align:center; font-weight:bold'>{row['Variação 1 ano (%)']:.1f}%</td>"
            html_setores += f"<td style='color:{cor_30d}; text-align:center; font-weight:bold'>{row['Variação 30 dias (%)']:.1f}%</td>"
            html_setores += "</tr>"

        html_setores += "</tbody></table>"

        st.markdown("### 🏙️ Variação Média por Setor")
        st.markdown(html_setores, unsafe_allow_html=True)

    def formatar_variacao(valor):
        cor = "green" if valor >= 1 else "red"
        return f"<td style='color:{cor}; text-align:center; font-weight:bold'>{valor:.1f}%</td>"

    def formatar_mayer(valor):
        if valor > 1.5:
            cor = "red"
        elif valor >= 1.0:
            cor = "orange"
        else:
            cor = "green"
        return f"<td style='color:{cor}; text-align:center; font-weight:bold'>{valor:.1f}</td>"

    def formatar_ticker(ticker):
        return f"<td style='text-align:center'><a href='?opcao=📈 Visão por Ativo&ticker={ticker}' target='_self'>{ticker}</a></td>"

    def formatar_td(texto):
        return f"<td style='text-align:center'>{texto}</td>"

    colunas = ["Ticker", "Nome da Empresa", "Variação 1 ano (%)", "Variação 30 dias (%)", "Índice Mayer", "Setor"]
    html = "<table style='width:100%; border-collapse: collapse;'>"
    html += "<thead><tr>" + "".join([f"<th style='text-align:center; border-bottom: 1px solid #ddd'>{col}</th>" for col in colunas]) + "</tr></thead><tbody>"

    for _, row in df_filtrado.iterrows():
        html += "<tr>"
        html += formatar_ticker(row["Ticker"])
        html += formatar_td(row["Nome da Empresa"])
        html += formatar_variacao(row["Variação 1 ano (%)"])
        html += formatar_variacao(row["Variação 30 dias (%)"])
        html += formatar_mayer(row["Índice Mayer"])
        html += formatar_td(row["Setor"])
        html += "</tr>"

    html += "</tbody></table>"

    st.markdown("### 📋 Tabela Consolidada dos Ativos")
    st.markdown(html, unsafe_allow_html=True)

# --- Análise Carteira ---
elif opcao == "📊 Análise Carteira":
    st.title("📊 Análise Carteira")

    df = carregar_dados()
    carteiras = carregar_carteiras_disponiveis()
    
    if not carteiras:
        st.warning("Nenhuma carteira encontrada. Por favor, crie uma carteira na aba '📂 Carteira de Investimento' primeiro.")
        st.stop()

    if carteiras:
        with st.container():
            st.markdown("#### Selecione a carteira:")
            col_sel, _ = st.columns([1, 3])
            with col_sel:
                carteira_escolhida = st.selectbox("", carteiras, key="selecao_carteira")

            df_carteira = carregar_carteira(carteira_escolhida)
            df_simulacao = simular_carteira(df, df_carteira.to_dict('records'))

            if not df_simulacao.empty:
                conn = sqlite3.connect(DB_PATH)
                df_fundamentais = pd.read_sql("SELECT ticker, nome, setor FROM dados_fundamentais", conn)
                df_indicadores = pd.read_sql("SELECT * FROM dados_indicadores", conn)
                conn.close()
                df_indicadores['data'] = pd.to_datetime(df_indicadores['data'])

                df_merge_setor = df_carteira.merge(df_fundamentais, on="ticker", how="left")
                df_merge_setor['aporte_monetario'] = df_merge_setor['quantidade'] * df_merge_setor['valor_compra']
                total_aporte = df_merge_setor['aporte_monetario'].sum()
                df_setores = df_merge_setor.groupby('setor')['aporte_monetario'].sum().reset_index()

                df_merge_setor['valor_total_aporte'] = df_merge_setor['quantidade'] * df_merge_setor['valor_compra']
                df_consolidado = df_merge_setor.groupby(['ticker', 'nome']).agg({
                    'quantidade': 'sum',
                    'valor_total_aporte': 'sum'
                }).reset_index()
                df_consolidado['preco_medio'] = df_consolidado['valor_total_aporte'] / df_consolidado['quantidade']

                desempenho_ativos = []
                for _, row in df_consolidado.iterrows():
                    ticker = row['ticker']
                    nome = row['nome']
                    qtd_total = row['quantidade']
                    preco_medio = row['preco_medio']

                    df_ativo = df[df['ticker'] == ticker].copy().sort_values('date')
                    if df_ativo.empty:
                        continue

                    preco_atual = df_ativo['close'].iloc[-1]
                    valor_atual = preco_atual * qtd_total
                    valor_aporte = preco_medio * qtd_total
                    ganho_monetario = valor_atual - valor_aporte
                    rendimento_pct = (preco_atual / preco_medio - 1) * 100
                    percentual_carteira = valor_aporte / total_aporte if total_aporte > 0 else 0

                    desempenho_ativos.append({
                        'Ticker': ticker,
                        'Nome da Empresa': nome,
                        'Quantidade': qtd_total,
                        'Preço Médio (R$)': f"R$ {preco_medio:.2f}",
                        '% da Carteira': percentual_carteira * 100,
                        'Rendimento (%)': rendimento_pct,
                        'Ganho/Perda (R$)': ganho_monetario,
                        'Valor Atual (R$)': valor_atual
                    })

                variacao_percentual_total = (df_simulacao['valor_total'].iloc[-1] / df_simulacao['aporte_acumulado'].iloc[-1] - 1) * 100
                ganho_total = df_simulacao['valor_total'].iloc[-1] - df_simulacao['aporte_acumulado'].iloc[-1]

                df_hist = df[df['ticker'].isin(df_carteira['ticker'].unique())]
                df_hist_pivot = df_hist.pivot(index='date', columns='ticker', values='close').dropna()
                df_retornos = df_hist_pivot.pct_change().dropna()
                pesos = df_merge_setor.groupby('ticker')['aporte_monetario'].sum()
                pesos /= pesos.sum()
                carteira_retornos = df_retornos @ pesos.loc[df_retornos.columns].values
                var_95 = np.percentile(carteira_retornos, 5) * 100

                correl_matrix = df_retornos.corr()
                correlacao_media = correl_matrix.where(np.tril(np.ones(correl_matrix.shape), k=-1).astype(bool)).stack().mean()

                cor_perf = "green" if variacao_percentual_total >= 0 else "red"
                cor_ganho = "green" if ganho_total >= 0 else "red"

                col1, col2, col3, col4 = st.columns(4)

                col1.markdown(f"""
                <div style='text-align:center;'>
                    <strong>📈 Performance da Carteira</strong><br>
                    <span style='font-size:26px; color:{cor_perf};'><strong>{variacao_percentual_total:.2f}%</strong></span>
                </div>
                """, unsafe_allow_html=True)

                col2.markdown(f"""
                <div style='text-align:center;'>
                    <strong>💰 Ganho/Perda Total (R$)</strong><br>
                    <span style='font-size:26px; color:{cor_ganho};'><strong>R$ {ganho_total:,.2f}</strong></span>
                </div>
                """, unsafe_allow_html=True)

                col3.markdown(f"""
                <div style='text-align:center;'>
                    <strong>🔻 VaR 95%</strong><br>
                    <span style='font-size:26px; color:#333333;'><strong>{var_95:.2f}%</strong></span>
                </div>
                """, unsafe_allow_html=True)

                col4.markdown(f"""
                <div style='text-align:center;'>
                    <strong>🔗 Correlação Média</strong><br>
                    <span style='font-size:26px; color:#333333;'><strong>{correlacao_media:.2f}</strong></span>
                </div>
                """, unsafe_allow_html=True)

                st.markdown("### 📈 Comparativo de Rentabilidade Acumulada")
                col_chk1, col_chk2, col_chk3 = st.columns(3)
                mostrar_ipca = col_chk1.checkbox("📈 Mostrar IPCA", value=True)
                mostrar_cdi = col_chk2.checkbox("📊 Mostrar CDI", value=True)
                mostrar_ibov = col_chk3.checkbox("📈 Mostrar IBOV", value=True)

                data_ini = df_simulacao['date'].min()
                data_fim = df_simulacao['date'].max()
                df_ind_periodo = df_indicadores[(df_indicadores['data'] >= data_ini) & (df_indicadores['data'] <= data_fim)].copy().sort_values('data')
                df_ind_periodo['ipca_acumulado'] = (1 + df_ind_periodo['ipca'].fillna(0)).cumprod() - 1
                df_ind_periodo['cdi_acumulado'] = (1 + df_ind_periodo['cdi'].fillna(0)).cumprod() - 1

                if 'ibov' in df_ind_periodo.columns and df_ind_periodo['ibov'].notna().sum() >= 2:
                    ibov_validos = df_ind_periodo['ibov'].dropna()
                    df_ind_periodo['ibov_acumulado'] = ibov_validos / ibov_validos.iloc[0] - 1
                else:
                    df_ind_periodo['ibov_acumulado'] = np.nan

                df_comp = df_simulacao[['date', 'rendimento_acumulado']].copy()
                df_comp = df_comp.merge(df_ind_periodo[['data', 'ipca_acumulado', 'cdi_acumulado', 'ibov_acumulado']],
                                        left_on='date', right_on='data', how='left').drop(columns='data')
                df_comp = df_comp.rename(columns={'rendimento_acumulado': 'Carteira'})

                df_plot = df_comp[['date', 'Carteira']]
                if mostrar_ipca:
                    df_plot['IPCA'] = df_comp['ipca_acumulado']
                if mostrar_cdi:
                    df_plot['CDI'] = df_comp['cdi_acumulado']
                if mostrar_ibov:
                    df_plot['IBOV'] = df_comp['ibov_acumulado']

                df_long = df_plot.melt(id_vars='date', var_name='Indicador', value_name='Valor')

                chart = alt.Chart(df_long).mark_line().encode(
                    x=alt.X('date:T', title='Data', axis=alt.Axis(format="%d/%m/%Y", labelAngle=-45)),
                    y=alt.Y('Valor:Q', title='Retorno Acumulado', axis=alt.Axis(format='%')),
                    color='Indicador:N',
                    tooltip=['date:T', 'Indicador:N', alt.Tooltip('Valor:Q', format='.2%')]
                ).properties(
                    width=800,
                    height=400,
                    title=f"Comparativo de Rentabilidade Acumulada - Carteira '{carteira_escolhida}'"
                ).interactive()
                st.altair_chart(chart, use_container_width=True)

                if not df_setores.empty:
                    st.markdown("### 🏅 Distribuição por Setor")
                    import plotly.express as px
                    fig = px.pie(
                        df_setores,
                        names="setor",
                        values="aporte_monetario",
                        title="Distribuição por Setor",
                        hole=0.4
                    )
                    fig.update_traces(
                        textposition='outside',
                        textinfo='percent+label',
                        pull=[0.05] * len(df_setores),
                        marker=dict(line=dict(color='#000000', width=1))
                    )
                    fig.update_layout(
                        showlegend=True,
                        height=420,
                        margin=dict(t=50, b=60, l=20, r=20)
                    )
                    st.plotly_chart(fig, use_container_width=True)

                df_exibicao = pd.DataFrame(desempenho_ativos)

                df_exibicao['% da Carteira'] = df_exibicao['% da Carteira'].astype(float).map(lambda x: f"{x:.2f}%")
                df_exibicao['Rendimento (%)'] = df_exibicao['Rendimento (%)'].astype(float)
                df_exibicao['Ganho/Perda (R$)'] = df_exibicao['Ganho/Perda (R$)'].astype(float).map(lambda x: f"R$ {x:,.2f}".replace(",", "v").replace(".", ",").replace("v", "."))
                df_exibicao['Valor Atual (R$)'] = df_exibicao['Valor Atual (R$)'].astype(float).map(lambda x: f"R$ {x:,.2f}".replace(",", "v").replace(".", ",").replace("v", "."))
                df_exibicao['Rendimento (%)'] = df_exibicao['Rendimento (%)'].map(lambda x: f"{x:.2f}%")

                gb = GridOptionsBuilder.from_dataframe(df_exibicao)
                gb.configure_default_column(filter=True, sortable=True, resizable=True, cellStyle={'textAlign': 'center'})

                js_code = JsCode("""
                function(params) {
                    if (params.colDef.field === 'Rendimento (%)' || params.colDef.field === 'Ganho/Perda (R$)') {
                        return {
                            'color': (parseFloat(params.value.replace(/[^0-9\-.,]/g, '').replace(',', '.')) < 0) ? 'red' : 'green',
                            'fontWeight': 'bold',
                            'textAlign': 'center'
                        }
                    } else {
                        return { 'textAlign': 'center' }
                    }
                }
                """)
                gb.configure_column("Rendimento (%)", cellStyle=js_code)
                gb.configure_column("Ganho/Perda (R$)", cellStyle=js_code)

                grid_options = gb.build()

                st.markdown("### 📋 Desempenho dos ativos na carteira (interativo)")
                AgGrid(
                    df_exibicao,
                    gridOptions=grid_options,
                    theme="material",
                    fit_columns_on_grid_load=True,
                    allow_unsafe_jscode=True
                )

# --- Backtesting HRP ---
if opcao == "🧪 Backtesting de Estratégias":
    st.title("🧪 Backtesting de Estratégias de Alocação")
    st.markdown("Compare diferentes estratégias de alocação de portfólio ao longo do tempo.")

    carteiras = carregar_carteiras_disponiveis()
    if not carteiras:
        st.warning("Nenhuma carteira encontrada. Crie uma carteira para usar esta funcionalidade.")
        st.stop()

    col1, col2 = st.columns(2)
    with col1:
        carteira_escolhida = st.selectbox("1. Universo de Ativos:", carteiras, key="backtest_carteira")
        estrategias_selecionadas = st.multiselect(
            "2. Estratégias para Simular:",
            options=["HRP", "Min Volatility", "Equal Weight", "Rebalancear para Pesos Iguais"],
            default=["HRP", "Equal Weight", "Rebalancear para Pesos Iguais"]
        )
    with col2:
        aporte_inicial = st.number_input("3. Aporte inicial (R$):", min_value=1000.0, value=10000.0, step=500.0)
        aporte_mensal = st.number_input("4. Aporte mensal (R$):", min_value=0.0, value=500.0, step=100.0)

    col3, col4 = st.columns(2)
    with col3:
        data_inicio_sim = st.date_input("5. Data de início da simulação:", value=datetime.today() - timedelta(days=365*10))
    with col4:
        max_aloc = st.number_input("6. Percentual máximo por ativo (%):", min_value=10.0, max_value=100.0, value=100.0, step=1.0)
    
    if 'resultados_backtest' not in st.session_state:
        st.session_state.resultados_backtest = None

    if st.button("🚀 Rodar Backtest Comparativo"):
        if not estrategias_selecionadas:
            st.error("Por favor, selecione pelo menos uma estratégia para simular.")
        else:
            with st.spinner("Executando simulações... Isso pode ser demorado."):
                df_carteira_selecionada = carregar_carteira(carteira_escolhida)
                tickers_simulacao = df_carteira_selecionada['ticker'].unique().tolist()
                data_fim_simulacao = df['date'].max()
                max_aloc_decimal = max_aloc / 100.0
                
                resultados = {}
                final_states = {}
                min_start_date = data_fim_simulacao

                for estrategia in estrategias_selecionadas:
                    st.write(f"Calculando estratégia: {estrategia}...")
                    df_resultado, final_state = run_backtest(
                        df, tickers_simulacao, pd.to_datetime(data_inicio_sim), data_fim_simulacao, 
                        aporte_inicial, aporte_mensal, max_aloc_decimal, estrategia
                    )
                    if not df_resultado.empty:
                        resultados[estrategia] = df_resultado
                        final_states[estrategia] = final_state
                        actual_start_date = final_state.get("actual_start_date")
                        if actual_start_date and actual_start_date < min_start_date:
                            min_start_date = actual_start_date
                
                st.session_state.resultados_backtest = {
                    "resultados": resultados, "final_states": final_states, "min_start_date": min_start_date,
                    "data_inicio_sim": data_inicio_sim, "data_fim_simulacao": data_fim_simulacao, "tickers_simulacao": tickers_simulacao
                }

    if st.session_state.resultados_backtest:
        resultados = st.session_state.resultados_backtest["resultados"]
        final_states = st.session_state.resultados_backtest["final_states"]
        min_start_date = st.session_state.resultados_backtest["min_start_date"]
        data_inicio_sim = st.session_state.resultados_backtest["data_inicio_sim"]
        data_fim_simulacao = st.session_state.resultados_backtest["data_fim_simulacao"]
        tickers_simulacao = st.session_state.resultados_backtest["tickers_simulacao"]

        if not resultados:
            st.error("Nenhuma estratégia pôde ser calculada. Verifique o período e os ativos.")
        else:
            st.markdown("---")
            if min_start_date and min_start_date.date() > data_inicio_sim:
                st.info(f"**Atenção:** As simulações começaram em **{min_start_date.strftime('%d/%m/%Y')}**. Esta foi a primeira data em que os critérios mínimos foram atendidos. Os benchmarks foram ajustados para esta data.")
            
            st.markdown("### 📈 Performance Comparativa das Estratégias")
            
            df_benchmarks = get_benchmark_data(min_start_date, data_fim_simulacao)
            df_plot_final = df_benchmarks.rename(columns={'data':'date'}).set_index('date')

            for estrategia, df_resultado in resultados.items():
                df_resultado = df_resultado.loc[min_start_date:]
                df_plot_final[estrategia] = (df_resultado['value'] / df_resultado['invested']) - 1
            
            df_plot_final.replace([np.inf, -np.inf], np.nan, inplace=True)
            df_plot_final.ffill(inplace=True)
            
            df_plot_final = df_plot_final.reset_index().rename(columns={'date':'data'})
            df_long = df_plot_final.melt(id_vars='data', var_name='Indicador', value_name='Valor')
            df_long['Valor'] = pd.to_numeric(df_long['Valor'], errors='coerce').fillna(0.0)

            chart = alt.Chart(df_long).mark_line().encode(
                x=alt.X('data:T', title='Data'),
                y=alt.Y('Valor:Q', title='Retorno Acumulado', axis=alt.Axis(format='%')),
                color='Indicador:N',
                tooltip=['data:T', 'Indicador:N', alt.Tooltip('Valor:Q', format='.2%')]
            ).properties(height=500).interactive()
            st.altair_chart(chart, use_container_width=True)

            st.markdown("### 📋 Painel de Performance das Estratégias")
            summary_data = []
            final_benchmarks_perf = { "IPCA": df_plot_final["IPCA"].iloc[-1], "CDI": df_plot_final["CDI"].iloc[-1], "IBOV": df_plot_final["IBOV"].iloc[-1] }
            
            def get_comparison_icon_html(strat_perf, bench_perf):
                return "🟢" if strat_perf > bench_perf else "🔴"

            for strat, df_result in resultados.items():
                valor_final = df_result['value'].iloc[-1]
                total_investido = df_result['invested'].iloc[-1]
                rendimento_final = (valor_final / total_investido - 1) if total_investido > 0 else 0
                summary_data.append({ "Estratégia": strat, "Patrimônio Final": valor_final, "Total Aportado": total_investido, "Rendimento Final": rendimento_final,
                    "vs IPCA": get_comparison_icon_html(rendimento_final, final_benchmarks_perf.get('IPCA', 0)),
                    "vs CDI": get_comparison_icon_html(rendimento_final, final_benchmarks_perf.get('CDI', 0)),
                    "vs IBOV": get_comparison_icon_html(rendimento_final, final_benchmarks_perf.get('IBOV', 0)) })
            df_summary_strat = pd.DataFrame(summary_data)
            st.dataframe(df_summary_strat.style.format({ "Patrimônio Final": "R$ {:,.2f}", "Total Aportado": "R$ {:,.2f}", "Rendimento Final": "{:.2%}" }), use_container_width=True)

            st.markdown("### 📊 Resumo Detalhado por Ativo")
            
            estrategia_para_resumo = st.selectbox("Selecione a estratégia para ver o resumo detalhado:", options=list(final_states.keys()) )

            if estrategia_para_resumo:
                final_state_resumo = final_states[estrategia_para_resumo]
                summary_list = []
                final_prices = final_state_resumo['prices']
                final_holdings = final_state_resumo['holdings']
                final_weights_ideal = final_state_resumo['weights']
                asset_starts = final_state_resumo['start_dates']
                total_invested_resumo = final_state_resumo.get('total_invested_per_asset', {})
                total_final_value = sum(final_holdings.get(t, 0) * final_prices.get(t, 0) for t in tickers_simulacao)

                df_hist_sim_period = df[(df['date'] >= pd.to_datetime(data_inicio_sim)) & (df['ticker'].isin(tickers_simulacao))]
                start_prices = df_hist_sim_period.groupby('ticker')['close'].first()
                
                for ticker in tickers_simulacao:
                    valor_final_ativo = final_holdings.get(ticker, 0) * final_prices.get(ticker, 0)
                    if valor_final_ativo == 0 and total_invested_resumo.get(ticker, 0) == 0:
                        continue 

                    preco_inicial_ativo = start_prices.get(ticker)
                    preco_final_ativo = final_prices.get(ticker)
                    perf_ativo = ((preco_final_ativo / preco_inicial_ativo) - 1) * 100 if preco_inicial_ativo and preco_final_ativo and preco_inicial_ativo > 0 else 0

                    summary_list.append({ "Ativo": ticker, "Data de Início na Carteira": asset_starts.get(ticker), "Total Aportado (R$)": total_invested_resumo.get(ticker, 0),
                        "Performance do Ativo (%)": perf_ativo, "Valor Final (R$)": valor_final_ativo,
                        "Proporção Final (%)": (valor_final_ativo / total_final_value) * 100 if total_final_value > 0 else 0,
                        "Proporção Ideal (%)": final_weights_ideal.get(ticker, 0) * 100 })
                
                df_summary = pd.DataFrame(summary_list).sort_values(by="Total Aportado (R$)", ascending=False)
                
                st.dataframe(df_summary.style.format({ "Data de Início na Carteira": lambda d: d.strftime('%d/%m/%Y') if pd.notnull(d) else '-',
                    "Total Aportado (R$)": "R$ {:,.2f}", "Performance do Ativo (%)": "{:.2f}%", "Valor Final (R$)": "R$ {:,.2f}",
                    "Proporção Final (%)": "{:.2f}%", "Proporção Ideal (%)": "{:.2f}%" }), use_container_width=True)

# --- Detecção de Rompimentos Técnicos ---
elif opcao == "📉 Detecção de Rompimentos Técnicos":
    st.title("📉 Detecção de Rompimentos Técnicos")

    janela = st.slider("Janela da média móvel (dias)", min_value=10, max_value=60, value=20)

    ultimo_dia = df['date'].max()
    st.markdown(f"Analisando dados mais recentes disponíveis: **{ultimo_dia.date()}**")

    resultados = []

    for ticker in sorted(df['ticker'].unique()):
        df_ativo = df[df['ticker'] == ticker].copy()
        df_ativo = df_ativo.sort_values("date")
        df_ativo['media_movel'] = df_ativo['close'].rolling(janela).mean()
        df_ativo['desvio'] = df_ativo['close'].rolling(janela).std()
        df_ativo['limite_sup'] = df_ativo['media_movel'] + 2 * df_ativo['desvio']
        df_ativo['limite_inf'] = df_ativo['media_movel'] - 2 * df_ativo['desvio']

        ultima_linha = df_ativo[df_ativo['date'] == ultimo_dia]
        if not ultima_linha.empty:
            preco = ultima_linha['close'].values[0]
            sup = ultima_linha['limite_sup'].values[0]
            inf = ultima_linha['limite_inf'].values[0]

            if preco > sup:
                rompimento = "Acima"
                sinal = "Aguardar"
            elif preco < inf:
                rompimento = "Abaixo"
                sinal = "Analisar compra"
            else:
                rompimento = "Dentro"
                sinal = "Aguardar"

            resultados.append({
                "Ticker": ticker,
                "Fechamento (R$)": f"R$ {preco:.2f}",
                "Rompimento": rompimento,
                "Sinal": sinal
            })

    df_resultado = pd.DataFrame(resultados)

    if not df_resultado.empty:
        st.markdown("### 📋 Status técnico no último dia disponível")
        st.dataframe(df_resultado, use_container_width=True)
    else:
        st.info("Não há dados suficientes para a análise.")

# --- Carteira de Investimento ---
elif opcao == "📂 Carteira de Investimento":
    st.title("📂 Carteira de Investimento")
    aba = st.radio("O que deseja fazer?", ["Criar nova carteira", "Importar carteira via Excel", "Visualizar carteira existente"])

    if aba == "Criar nova carteira":
        st.subheader("Criar uma nova carteira")
        nome_carteira = st.text_input("Nome da carteira")
        num_linhas = st.number_input("Quantos ativos deseja inserir?", min_value=1, max_value=20, value=3)

        dados = []
        for i in range(num_linhas):
            st.markdown(f"### Ativo {i+1}")
            cols = st.columns(4)
            ticker = cols[0].text_input(f"Ticker {i+1}", key=f"ticker_{i}")
            quantidade = cols[1].number_input(f"Quantidade {i+1}", min_value=1, key=f"qtd_{i}")
            valor = cols[2].number_input(f"Valor de compra (R$) {i+1}", min_value=0.0, step=0.01, key=f"valor_{i}")
            data = cols[3].date_input(f"Data da compra {i+1}", key=f"data_{i}")
            dados.append({"ticker": ticker, "quantidade": quantidade, "valor_compra": valor, "data_compra": str(data)})

        if st.button("Salvar carteira"):
            if nome_carteira.strip() == "":
                st.error("Digite um nome para a carteira.")
            else:
                salvar_na_carteira(nome_carteira.strip(), dados)
                st.success(f"Carteira '{nome_carteira}' salva com sucesso!")

    elif aba == "Importar carteira via Excel":
        st.subheader("📥 Importar carteira via Excel (usando tabela nomeada)")
        uploaded_file = st.file_uploader("Selecione o arquivo Excel com a carteira (com tabela nomeada)", type=["xlsx"])
        nome_importacao = st.text_input("Nome para a carteira importada", key="nome_import")

        if st.button("Importar dados do Excel"):
            if not uploaded_file:
                st.error("Por favor, selecione um arquivo Excel.")
            elif nome_importacao.strip() == "":
                st.error("Dê um nome para essa carteira importada.")
            else:
                try:
                    df_excel = importar_tabela_excel(uploaded_file)
                    colunas_esperadas = {'data', 'ticker', 'quantidade', 'valor_compra'}
                    if not colunas_esperadas.issubset(set(df_excel.columns)):
                        st.error(f"O arquivo deve conter as colunas: {', '.join(colunas_esperadas)}.")
                    else:
                        dados_convertidos = []
                        for _, row in df_excel.iterrows():
                            dados_convertidos.append({
                                "ticker": row['ticker'],
                                "quantidade": int(row['quantidade']),
                                "valor_compra": float(row['valor_compra']),
                                "data_compra": str(pd.to_datetime(row['data']).date())
                            })

                        conn = sqlite3.connect(DB_PATH)
                        conn.execute("DELETE FROM carteiras WHERE nome = ?", (nome_importacao.strip(),))
                        conn.commit()
                        conn.close()

                        salvar_na_carteira(nome_importacao.strip(), dados_convertidos)
                        st.success(f"Carteira '{nome_importacao}' importada e sobrescrita com sucesso!")
                except Exception as e:
                    st.error(f"Erro ao importar arquivo: {e}")

    elif aba == "Visualizar carteira existente":
        st.subheader("Consultar e editar carteira")
        carteiras = carregar_carteiras_disponiveis()
        if carteiras:
            carteira_escolhida = st.selectbox("Selecione a carteira:", carteiras)

            if st.button("🗑️ Deletar esta carteira"):
                try:
                    conn = sqlite3.connect(DB_PATH)
                    conn.execute("DELETE FROM carteiras WHERE nome = ?", (carteira_escolhida,))
                    conn.commit()
                    conn.close()
                    st.success(f"Carteira '{carteira_escolhida}' deletada com sucesso.")
                    try:
                        st.rerun()
                    except AttributeError:
                        st.experimental_rerun()
                except Exception as e:
                    st.error(f"Erro ao deletar carteira: {e}")
                st.stop()

            df_carteira = carregar_carteira(carteira_escolhida)
            st.dataframe(df_carteira)

            st.markdown("---")
            st.markdown("### Editar ou adicionar novos ativos à carteira")

            for i, row in df_carteira.iterrows():
                cols = st.columns(4)
                qtd = cols[0].number_input(f"Quantidade ({row['ticker']})", min_value=1, value=row['quantidade'], key=f"edit_qtd_{row['id']}")
                val = cols[1].number_input(f"Valor compra ({row['ticker']})", min_value=0.0, value=row['valor_compra'], step=0.01, key=f"edit_val_{row['id']}")
                data = cols[2].date_input(f"Data compra ({row['ticker']})", value=pd.to_datetime(row['data_compra']), key=f"edit_data_{row['id']}")
                if cols[3].button("Salvar", key=f"salvar_{row['id']}"):
                    atualizar_entrada_carteira(row['id'], qtd, val, str(data))
                    st.success(f"Ativo {row['ticker']} atualizado com sucesso!")
        else:
            st.info("Nenhuma carteira cadastrada ainda. Crie uma nova carteira na aba anterior.")

# --- Otimizador Avançado ---
elif opcao == "🤖 Otimizador Avançado (HRP)":
    st.title("🤖 Otimizador Avançado de Risco (HRP)")
    st.markdown("""
    Esta abordagem utiliza **Hierarchical Risk Parity (HRP)** para criar uma carteira diversificada com base no risco e na correlação dos ativos. 
    Abaixo, você pode opcionalmente aplicar um limite máximo de alocação por ativo.
    """)

    carteiras = carregar_carteiras_disponiveis()
    if not carteiras:
        st.warning("Nenhuma carteira encontrada. Por favor, crie uma carteira na aba '📂 Carteira de Investimento' primeiro.")
        st.stop()

    col1, col2 = st.columns(2)
    with col1:
        carteira_escolhida = st.selectbox("Selecione a carteira para otimizar:", carteiras, key="hrp_carteira")
    with col2:
        valor_aporte = st.number_input("Valor do novo aporte (R$):", min_value=1.0, value=1000.0, step=100.0, key="hrp_aporte")

    st.markdown("#### 🛡️ Controle de Risco e Diversificação (Opcional)")
    max_aloc = st.number_input(
        "Defina o percentual máximo por ativo (ex: 25 para 25%):",
        min_value=10.0,
        max_value=100.0,
        value=100.0,
        step=1.0
    )
    max_aloc_decimal = max_aloc / 100.0

    if st.button("🧠 Calcular Otimização HRP"):
        with st.spinner("Analisando correlações, otimizando e aplicando restrições..."):
            try:
                df_carteira = carregar_carteira(carteira_escolhida)
                tickers = df_carteira['ticker'].unique().tolist()
                df_historico_completo = carregar_dados()
                df_precos = df_historico_completo[df_historico_completo['ticker'].isin(tickers)]
                df_precos_pivot = df_precos.pivot(index='date', columns='ticker', values='close').dropna()
                retornos = df_precos_pivot.pct_change().dropna()

                hrp = HRPOpt(retornos)
                hrp.optimize()
                pesos_hrp_puros = hrp.clean_weights()

                if max_aloc_decimal < 1.0:
                    pesos_otimizados = rebalance_weights_with_cap(pesos_hrp_puros, max_aloc_decimal)
                else:
                    pesos_otimizados = pesos_hrp_puros

                st.markdown("### 🎯 Alocação Otimizada por Paridade de Risco (com Restrições)")
                
                fig, ax = plt.subplots(figsize=(10, 6))
                plotting.plot_dendrogram(hrp, ax=ax)
                st.pyplot(fig)
                
                st.markdown("### 💸 Sugestão de Alocação do Aporte")

                precos_recentes = get_latest_prices(df_precos_pivot)
                df_carteira['valor_atual'] = df_carteira.apply(
                    lambda row: row['quantidade'] * precos_recentes.get(row['ticker'], 0), axis=1
                )
                valor_carteira_atual = df_carteira['valor_atual'].sum()
                valor_total_novo = valor_carteira_atual + valor_aporte

                resultados = []
                for ticker, peso in pesos_otimizados.items():
                    valor_atual_ativo = df_carteira[df_carteira['ticker'] == ticker]['valor_atual'].sum()
                    valor_otimizado_ativo = valor_total_novo * peso
                    aporte_sugerido = max(0, valor_otimizado_ativo - valor_atual_ativo)
                    
                    resultados.append({
                        "Ativo": ticker,
                        "Alocação Atual (%)": (valor_atual_ativo / valor_carteira_atual * 100) if valor_carteira_atual > 0 else 0,
                        "Alocação Otimizada HRP (%)": peso * 100,
                        "Aporte Sugerido (R$)": aporte_sugerido
                    })
                
                df_resultados = pd.DataFrame(resultados).sort_values(by="Alocação Otimizada HRP (%)", ascending=False)
                
                total_sugerido = df_resultados['Aporte Sugerido (R$)'].sum()
                if total_sugerido > 0:
                    df_resultados['Aporte Sugerido (R$)'] = (df_resultados['Aporte Sugerido (R$)'] / total_sugerido) * valor_aporte
                
                df_resultados['Ações a Comprar'] = df_resultados.apply(
                    lambda row: int(row['Aporte Sugerido (R$)'] / precos_recentes.get(row['Ativo'])) if precos_recentes.get(row['Ativo'], 0) > 0 else 0,
                    axis=1
                )

                ultimo_dia_df = df_historico_completo['date'].max()
                tickers_carteira = df_resultados['Ativo'].unique().tolist()
                
                performances = calcular_performances(df_historico_completo, tickers_carteira, ultimo_dia_df)

                df_resultados['Performance 30d (%)'] = df_resultados['Ativo'].map(lambda t: performances.get(t, {}).get('30d', 0) * 100)
                df_resultados['Performance 1 Ano (%)'] = df_resultados['Ativo'].map(lambda t: performances.get(t, {}).get('1y', 0) * 100)
                df_resultados['Performance YTD (%)'] = df_resultados['Ativo'].map(lambda t: performances.get(t, {}).get('ytd', 0) * 100)
                
                colunas_display = [
                    "Ativo", 'Performance 30d (%)', 'Performance 1 Ano (%)', 'Performance YTD (%)',
                    "Alocação Atual (%)", "Alocação Otimizada HRP (%)", "Aporte Sugerido (R$)", "Ações a Comprar"
                ]
                df_display = df_resultados[colunas_display]

                colunas_html_header = [
                    "Ativo", "Perf. 30D", "Perf. 1 Ano", "Perf. YTD",
                    "Aloc. Atual (%)", "Aloc. HRP (%)", "Aporte Sugerido (R$)", "Ações a Comprar"
                ]

                html_table = "<table style='width:100%; border-collapse: collapse;'>"
                html_table += "<thead><tr>" + "".join([f"<th style='text-align:center; border-bottom: 1px solid #ddd'>{col}</th>" for col in colunas_html_header]) + "</tr></thead><tbody>"

                for _, row in df_display.iterrows():
                    link_ativo = f"<a href='?opcao=📈 Visão por Ativo&ticker={row['Ativo']}' target='_self'>{row['Ativo']}</a>"
                    cor_30d = "green" if row['Performance 30d (%)'] >= 0 else "red"
                    cor_1y = "green" if row['Performance 1 Ano (%)'] >= 0 else "red"
                    cor_ytd = "green" if row['Performance YTD (%)'] >= 0 else "red"

                    html_table += "<tr>"
                    html_table += f"<td style='text-align:center'>{link_ativo}</td>"
                    html_table += f"<td style='color:{cor_30d}; text-align:center; font-weight:bold'>{row['Performance 30d (%)']:.2f}%</td>"
                    html_table += f"<td style='color:{cor_1y}; text-align:center; font-weight:bold'>{row['Performance 1 Ano (%)']:.2f}%</td>"
                    html_table += f"<td style='color:{cor_ytd}; text-align:center; font-weight:bold'>{row['Performance YTD (%)']:.2f}%</td>"
                    html_table += f"<td style='text-align:center'>{row['Alocação Atual (%)']:.2f}%</td>"
                    html_table += f"<td style='text-align:center; font-weight:bold'>{row['Alocação Otimizada HRP (%)']:.2f}%</td>"
                    html_table += f"<td style='text-align:center'>R$ {row['Aporte Sugerido (R$)']:.2f}</td>"
                    html_table += f"<td style='text-align:center'>{row['Ações a Comprar']}</td>"
                    html_table += "</tr>"

                html_table += "</tbody></table>"

                st.markdown(html_table, unsafe_allow_html=True)
                
                col_graf1, col_graf2 = st.columns(2)

                with col_graf1:
                    fig_aloc = go.Figure(data=[go.Pie(
                        labels=df_resultados['Ativo'], 
                        values=df_resultados['Alocação Otimizada HRP (%)'], 
                        hole=.3,
                        textinfo='percent+label'
                    )])
                    fig_aloc.update_layout(
                        title_text='Alocação Otimizada da Carteira Total',
                        title_x=0.5,
                        title_xanchor='center'
                    )
                    st.plotly_chart(fig_aloc, use_container_width=True)
                
                with col_graf2:
                    df_aporte_plot = df_resultados[df_resultados['Aporte Sugerido (R$)'] > 0.01]
                    fig_aporte = go.Figure(data=[go.Pie(
                        labels=df_aporte_plot['Ativo'], 
                        values=df_aporte_plot['Aporte Sugerido (R$)'], 
                        hole=.3,
                        textinfo='percent+label'
                    )])
                    fig_aporte.update_layout(
                        title_text=f'Distribuição do Aporte de R$ {valor_aporte:,.2f}',
                        title_x=0.5,
                        title_xanchor='center'
                    )

                    if not df_aporte_plot.empty:
                        st.plotly_chart(fig_aporte, use_container_width=True)
                    else:
                        st.info("Nenhum aporte sugerido para estes ativos. A carteira já está próxima da alocação de risco ideal.")
                
                valor_sobra = valor_aporte - (df_resultados['Ações a Comprar'] * df_resultados['Ativo'].map(precos_recentes)).sum()
                st.info(f"Valor restante do aporte (não alocado para compra de ações inteiras): **R$ {valor_sobra:.2f}**")

            except Exception as e:
                st.error(f"Ocorreu um erro durante a otimização HRP: {e}")
                st.error("Verifique se há dados históricos suficientes para todos os ativos na carteira.")